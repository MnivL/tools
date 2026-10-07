"""Excel 写出公共模块。

封装 openpyxl 的写入逻辑，提供：
- 表头字段推断
- 嵌套对象拍平
- 基本类型安全序列化
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
import pickle
import re
import tempfile
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
DEFAULT_MAX_SHEETS_PER_WORKBOOK = 50


def flatten_dict(data: Mapping[str, Any], parent_key: str = "", sep: str = ".") -> dict[str, Any]:
    """将嵌套字典拍平为单层字典。

    例如：``{"a": {"b": 1}}`` -> ``{"a.b": 1}``。
    列表类型的值不递归展开，保留为整体。
    """

    items: dict[str, Any] = {}
    for key, value in data.items():
        new_key = f"{parent_key}{sep}{key}" if parent_key else str(key)
        if isinstance(value, Mapping):
            items.update(flatten_dict(value, new_key, sep=sep))
        else:
            items[new_key] = value
    return items


def collect_columns(rows: Iterable[Mapping[str, Any]]) -> list[str]:
    """汇总所有行中出现的字段，按首次出现顺序返回作为表头。"""

    seen: dict[str, None] = {}
    for row in rows:
        for key in flatten_dict(row).keys():
            if key not in seen:
                seen[key] = None
    return list(seen.keys())


def serialize_value(value: Any) -> Any:
    """将无法直接写入 Excel 的值转换为兼容类型。"""

    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    # 其他不可识别的对象统一使用 JSON 字符串表达
    import json

    try:
        return json.dumps(value, ensure_ascii=False, default=str)
    except Exception:
        return str(value)


def write_rows_to_sheet(worksheet: Worksheet, headers: Sequence[str], data_rows: Sequence[Mapping[str, Any]]) -> None:
    """将数据按表头顺序写入指定工作表。"""

    worksheet.append(list(headers))
    for row in data_rows:
        flat = flatten_dict(row)
        worksheet.append([serialize_value(flat.get(col)) for col in headers])


def _column_groups(headers: Sequence[str], max_columns: int) -> list[Sequence[str]]:
    if max_columns < 1 or max_columns > EXCEL_MAX_COLUMNS:
        raise ValueError(f"max_columns 必须在 1 到 {EXCEL_MAX_COLUMNS} 之间")
    return [headers[start : start + max_columns] for start in range(0, len(headers), max_columns)] or [headers]


def _safe_sheet_title(name: str, used: set[str], index: int) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "data"
    candidate = cleaned[:31]
    if candidate not in used:
        used.add(candidate)
        return candidate
    suffix = f"_{index}"
    candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
    while candidate in used:
        index += 1
        suffix = f"_{index}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
    used.add(candidate)
    return candidate


def _write_partitioned_rows(
    workbook: Workbook,
    rows: Iterable[Mapping[str, Any]],
    headers: Sequence[str],
    sheet_name: str,
    max_rows: int = EXCEL_MAX_ROWS,
    max_columns: int = EXCEL_MAX_COLUMNS,
) -> None:
    if max_rows < 2 or max_rows > EXCEL_MAX_ROWS:
        raise ValueError(f"max_rows 必须在 2 到 {EXCEL_MAX_ROWS} 之间")
    groups = _column_groups(headers, max_columns)
    used: set[str] = set()
    worksheets: list[Worksheet] = []
    for group_index, group in enumerate(groups, start=1):
        title = sheet_name if len(groups) == 1 else f"{sheet_name}_c{group_index}"
        worksheet = workbook.create_sheet(_safe_sheet_title(title, used, group_index))
        worksheet.append(list(group))
        worksheets.append(worksheet)

    data_count = 0
    row_group = 1
    for row in rows:
        if data_count >= max_rows - 1:
            data_count = 0
            row_group += 1
            for group_index, group in enumerate(groups, start=1):
                title = f"{sheet_name}_r{row_group}" if len(groups) == 1 else f"{sheet_name}_r{row_group}_c{group_index}"
                worksheet = workbook.create_sheet(_safe_sheet_title(title, used, row_group * 1000 + group_index))
                worksheet.append(list(group))
                worksheets[group_index - 1] = worksheet
        flat = flatten_dict(row)
        for worksheet, group in zip(worksheets, groups):
            worksheet.append([serialize_value(flat.get(col)) for col in group])
        data_count += 1


def auto_adjust_width(worksheet: Worksheet, max_width: int = 60) -> None:
    """根据内容自动调整列宽，提升可读性。"""

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        longest = 0
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            length = len(str(value))
            if length > longest:
                longest = length
        worksheet.column_dimensions[column_letter].width = min(max(longest + 2, 10), max_width)


def save_workbook(workbook: Workbook, output_path: str | Path) -> Path:
    """保存工作簿并以 .xlsx 后缀规范路径。"""

    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def rows_to_workbook(
    rows: Sequence[Mapping[str, Any]],
    sheet_name: str = "data",
    max_rows: int = EXCEL_MAX_ROWS,
    max_columns: int = EXCEL_MAX_COLUMNS,
) -> Workbook:
    """将行数据构建为 Workbook 对象，便于二次处理。"""

    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is not None:
        workbook.remove(worksheet)
    headers = collect_columns(rows)
    _write_partitioned_rows(workbook, rows, headers, sheet_name, max_rows, max_columns)
    for worksheet in workbook.worksheets:
        auto_adjust_width(worksheet)
    return workbook


def stream_rows_to_files(
    rows: Iterable[Mapping[str, Any]],
    output_path: str | Path,
    sheet_name: str = "data",
    max_rows: int = EXCEL_MAX_ROWS,
    max_columns: int = EXCEL_MAX_COLUMNS,
    max_sheets_per_workbook: int = DEFAULT_MAX_SHEETS_PER_WORKBOOK,
) -> list[Path]:
    """流式写入 Excel，超出工作表/工作簿容量时自动分片。

    行数据先落到临时磁盘文件，以便在不把全部数据放进内存的情况下先确定完整表头。
    返回生成的一个或多个文件路径；单文件时路径就是传入的输出路径。
    """
    if max_sheets_per_workbook < 1:
        raise ValueError("max_sheets_per_workbook 必须大于 0")
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        output = output.with_suffix(".xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)

    headers: list[str] = []
    seen: set[str] = set()
    with tempfile.TemporaryFile(mode="w+b") as spool:
        for row in rows:
            if not isinstance(row, Mapping):
                raise ValueError("每一行数据必须是对象")
            materialized = dict(row)
            pickle.dump(materialized, spool, protocol=pickle.HIGHEST_PROTOCOL)
            for key in flatten_dict(materialized):
                if key not in seen:
                    seen.add(key)
                    headers.append(key)

        groups = _column_groups(headers, max_columns)
        data_capacity = max_rows - 1
        if data_capacity < 1:
            raise ValueError(f"max_rows 必须在 2 到 {EXCEL_MAX_ROWS} 之间")

        output_paths: list[Path] = []
        workbook: Workbook | None = None
        worksheets: list[Any] = []
        widths: list[list[int]] = []
        row_in_sheet = data_capacity
        row_group = 0
        workbook_sheet_count = 0
        workbook_index = 0

        def start_workbook() -> None:
            nonlocal workbook, worksheets, widths, row_in_sheet, row_group, workbook_sheet_count
            workbook = Workbook(write_only=True)
            worksheets = []
            widths = []
            row_in_sheet = data_capacity
            row_group = 0
            workbook_sheet_count = 0

        def start_row_group() -> None:
            nonlocal row_in_sheet, row_group, workbook_sheet_count
            assert workbook is not None
            row_group += 1
            row_in_sheet = 0
            worksheets.clear()
            widths.clear()
            for group_index, group in enumerate(groups, start=1):
                title = sheet_name if row_group == 1 and len(groups) == 1 else f"{sheet_name}_r{row_group}_c{group_index}"
                worksheet = workbook.create_sheet(_safe_sheet_title(title, set(), row_group * 1000 + group_index))
                worksheet.append(list(group))
                worksheets.append(worksheet)
                widths.append([min(max(len(str(value)) + 2, 10), 60) for value in group])
                workbook_sheet_count += 1

        def save_current() -> None:
            nonlocal workbook, workbook_index
            if workbook is None:
                return
            for worksheet, sheet_widths in zip(worksheets, widths):
                for index, width in enumerate(sheet_widths, start=1):
                    worksheet.column_dimensions[get_column_letter(index)].width = width
            workbook_index += 1
            path = output if workbook_index == 1 else output.with_name(f"{output.stem}.part{workbook_index:03d}{output.suffix}")
            workbook.save(path)
            output_paths.append(path)
            workbook = None

        def iter_spooled_rows() -> Iterator[Mapping[str, Any]]:
            while True:
                try:
                    yield pickle.load(spool)
                except EOFError:
                    return

        start_workbook()
        spool.seek(0)
        for row in iter_spooled_rows():
            if workbook is None:
                start_workbook()
            if row_in_sheet >= data_capacity:
                if workbook_sheet_count + len(groups) > max_sheets_per_workbook and workbook_sheet_count:
                    save_current()
                    start_workbook()
                start_row_group()
            flat = flatten_dict(row)
            for worksheet, group, sheet_widths in zip(worksheets, groups, widths):
                values = [serialize_value(flat.get(col)) for col in group]
                worksheet.append(values)
                for index, value in enumerate(values):
                    if value is not None:
                        sheet_widths[index] = min(max(sheet_widths[index], len(str(value)) + 2), 60)
            row_in_sheet += 1
        if not headers:
            start_row_group()
        save_current()

    return output_paths or [output]
