import json
from io import StringIO
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from any2excel.excel_writer import flatten_dict, serialize_value, stream_rows_to_files
from any2excel.json2excel import convert as convert_json, iter_json_array_stream
from any2excel.jsonl2excel import iter_jsonl


class Any2ExcelTests(unittest.TestCase):
    def test_flatten_and_serialize_special_values(self):
        self.assertEqual(flatten_dict({"a": {"b": 1}, "items": [1, 2]}), {"a.b": 1, "items": [1, 2]})
        self.assertEqual(serialize_value(date(2026, 10, 6)), "2026-10-06")
        self.assertEqual(serialize_value(Decimal("1.20")), "1.20")
        self.assertEqual(serialize_value({"name": "张三"}), '{"name": "张三"}')

    def test_json_array_conversion_round_trip(self):
        rows = [
            {"id": 1, "profile": {"city": "上海"}},
            {"id": 2, "name": "李四", "tags": ["a", "b"]},
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = convert_json(rows, Path(directory) / "result")
            self.assertEqual(output.suffix, ".xlsx")
            values = list(load_workbook(output).active.values)
        self.assertEqual(values[0], ("id", "profile.city", "name", "tags"))
        self.assertEqual(values[1], (1, "上海", None, None))
        self.assertEqual(values[2], (2, None, "李四", '["a", "b"]'))

    def test_jsonl_parser_skips_comments_and_reports_line_number(self):
        rows = list(iter_jsonl('# comment\n\n{"id": 1}\n'))
        self.assertEqual(rows, [{"id": 1}])
        with self.assertRaisesRegex(ValueError, r"第 2 行 JSON 解析失败"):
            list(iter_jsonl('{"id": 1}\nnot-json\n'))

    def test_jsonl_parser_rejects_non_object(self):
        with self.assertRaisesRegex(ValueError, r"第 1 行不是 JSON 对象"):
            list(iter_jsonl("[1, 2]\n"))

    def test_streaming_json_array_handles_small_chunks(self):
        stream = StringIO('[{"id": 1, "name": "a"}, {"id": 2, "name": "b"}]')
        self.assertEqual(list(iter_json_array_stream(stream, chunk_size=3)), [{"id": 1, "name": "a"}, {"id": 2, "name": "b"}])
        self.assertEqual(list(iter_json_array_stream(StringIO("  [ ]  "), chunk_size=1)), [])

    def test_streaming_writer_splits_rows_and_workbooks(self):
        rows = ({"id": index, "long": "x"} for index in range(5))
        with tempfile.TemporaryDirectory() as directory:
            paths = stream_rows_to_files(
                rows,
                Path(directory) / "large-output.xlsx",
                max_rows=3,
                max_sheets_per_workbook=1,
            )
            self.assertEqual([path.name for path in paths], ["large-output.xlsx", "large-output.part002.xlsx", "large-output.part003.xlsx"])
            self.assertEqual(list(load_workbook(paths[0]).active.values), [("id", "long"), (0, "x"), (1, "x")])
            self.assertEqual(list(load_workbook(paths[2]).active.values), [("id", "long"), (4, "x")])


if __name__ == "__main__":
    unittest.main()
